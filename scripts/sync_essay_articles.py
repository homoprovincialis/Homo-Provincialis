#!/usr/bin/env python3
"""Synchronize readable articles, Word downloads and the bilingual essay index.

Usage: python scripts/sync_essay_articles.py

Source documents and generated files live together in homoprovincialis-site/Articles.
Existing originals are not overwritten. Newly added articles are discovered without
editing a hard-coded list. The branch workflow commits generated files to the
feature branch for review; it never commits directly to main.
"""
from __future__ import annotations

import re
from html import escape
from pathlib import Path

import mammoth
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SITE = ROOT / "homoprovincialis-site"
ARTICLES = SITE / "Articles"
INDEX = SITE / "essays-by-ramon-eyquem.html"

WORD_PAGE_CSS = """
*{box-sizing:border-box}
body{margin:0;background:#eef1f3;color:#28333d;font:17px/1.7 Georgia,'Times New Roman',serif}
.topbar{background:#fff;border-bottom:1px solid #d7dfe5;padding:13px max(18px,calc((100vw - 920px)/2));font:14px Arial,sans-serif}
a{color:#245779;overflow-wrap:anywhere}
.article{max-width:920px;margin:28px auto 50px;padding:clamp(22px,5vw,60px);background:#fff;border:1px solid #d5dde3}
h1,h2,h3{font-family:Georgia,'Times New Roman',serif;color:#23445c;line-height:1.35}
h1{font-size:clamp(27px,4vw,36px);font-weight:normal}
h2{font-size:23px}h3{font-size:20px}p{margin:.4em 0 .95em}
.article{overflow-wrap:anywhere}.article img{max-width:100%;height:auto}
.article table{display:block;overflow-x:auto;border-collapse:collapse}
.article td,.article th{border:1px solid #c9d3da;padding:7px 10px}
.download{font:14px Arial,sans-serif;margin:30px 0 0;padding-top:17px;border-top:1px solid #d7dfe5}
@media(max-width:650px){body{background:#fff;font-size:16px}.article{margin:0;padding:18px;border:0}}
"""


def docx_to_html(source: Path, target: Path) -> None:
    with source.open("rb") as handle:
        result = mammoth.convert_to_html(handle)
    if not result.value.strip():
        raise ValueError("Word conversion produced an empty article: " + source.name)
    suffix = source.stem.lower()
    lang = "ru" if suffix.endswith("_ru") else "en"
    title = source.stem.replace("_", " ")
    doc = Document(source)
    if doc.core_properties.title and doc.core_properties.title.strip():
        title = doc.core_properties.title.strip()
    else:
        for p in doc.paragraphs:
            if p.text.strip():
                title = p.text.strip()
                break
    back = "← Эссе Рамона Эйкема" if lang == "ru" else "← Essays by Ramon Eyquem"
    download = "Скачать документ Word" if lang == "ru" else "Download Word document"
    html = f"""<!doctype html><html lang="{lang}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)} — Homo Provincialis</title><style>{WORD_PAGE_CSS}</style></head>
<body><nav class="topbar"><a href="../essays-by-ramon-eyquem.html">{back}</a></nav>
<main class="article"><h1>{escape(title)}</h1><article>{result.value}</article>
<p class="download"><a href="{escape(source.name, quote=True)}" download>{download} ↓</a></p>
</main></body></html>"""
    target.write_text(html, encoding="utf-8")
    print("Created browser article:", target.name)


def html_to_docx(source: Path, target: Path) -> None:
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    main = soup.find("main")
    if main is None:
        raise ValueError("Article HTML lacks a main element: " + source.name)
    lines = []
    for tag in main.find_all(["section", "h1", "h2", "h3", "p", "li"]):
        if tag.name == "section":
            if "lang-section" in tag.get("class", []):
                code = tag.get("lang") or tag.get("id", "")
                lines.append(("Heading 1", {"en":"English","ru":"Русский","fr":"Français","ro":"Română"}.get(code, code)))
            continue
        if tag.find_parent(["nav", "script", "style"]):
            continue
        if "download" in tag.get("class", []):
            continue
        if tag.name == "li" and tag.find_parent("li"):
            continue
        text = tag.get_text(" ", strip=True)
        if text:
            lines.append(({"h1":"Title","h2":"Heading 2","h3":"Heading 3","p":"Normal","li":"List Bullet"}[tag.name], text))
    if len(lines) < 3 or sum(len(text) for _, text in lines) < 160:
        raise ValueError("Refusing to create an incomplete Word document: " + source.name)
    doc = Document()
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(2.25), Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.55)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)
    base.paragraph_format.space_after, base.paragraph_format.line_spacing = Pt(7), 1.16
    for heading in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[heading].font.name = "Cambria"
        doc.styles[heading].font.color.rgb = RGBColor(34, 68, 93)
    for style, text in lines:
        doc.add_paragraph(text, style=style)
    doc.core_properties.title = lines[0][1]
    doc.core_properties.author = "Ramon Eyquem"
    doc.save(target)
    print("Created Word download:", target.name)


def languages_for_article(source: Path) -> set[str]:
    stem = source.stem.lower()
    if stem.endswith("_ro_en_fr_ru"):
        return {"en", "ru"}
    if stem.endswith("_ru") or stem.endswith("_ru_ru"):
        return {"ru"}
    if stem.endswith("_en"):
        return {"en"}
    # Distinguish original Romanian/French texts from the two visible site languages.
    if stem.endswith(("_ro", "_fr")):
        return set()
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    lang = (soup.html.get("lang", "") if soup.html else "").lower()
    if lang in ("en", "ru"):
        return {lang}
    return set()


def title_for_article(source: Path, language: str) -> str:
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    section = soup.find("section", class_="lang-section", lang=language)
    h1 = (section.find("h1") if section else None) or soup.find("main").find("h1") if soup.find("main") else None
    if h1 and h1.get_text(strip=True):
        return h1.get_text(" ", strip=True)
    if soup.title:
        return soup.title.get_text(" ", strip=True).replace(" — Homo Provincialis", "")
    return source.stem.replace("_", " ")


def sync_listing(content: str, language: str, files: list[Path]) -> str:
    label = "en" if language == "en" else "ru"
    section_start = content.find(f'<section class="lang lang-{label} essay-language"')
    if section_start < 0:
        raise ValueError("Missing essay language section: " + label)
    section_end = content.find("</section>", section_start)
    if section_end < 0:
        raise ValueError("Incomplete essay language section: " + label)
    before, section, after = content[:section_start], content[section_start:section_end], content[section_end:]
    used = set(re.findall(r'href="Articles/([^"#]+)\.html(?:#[a-z]+)?"', section))
    new_entries = []
    for file in files:
        if language not in languages_for_article(file) or file.stem in used:
            continue
        title = title_for_article(file, language)
        safe = escape(file.stem, quote=True)
        href = f'Articles/{safe}.html' + ("#" + language if file.stem.endswith("_RO_EN_FR_RU") else "")
        category = "Article" if language == "en" else "Статья"
        description = ("Available to read online or download as a Word document."
                       if language == "en" else
                       "Доступна для чтения на сайте и скачивания в формате Word.")
        read = "Read article" if language == "en" else "Читать статью"
        download = "Download" if language == "en" else "Скачать"
        new_entries.append(f'''        <article class="essay-entry">
          <div class="essay-number" aria-hidden="true">00</div>
          <div class="essay-content">
            <div class="essay-category">{category}</div>
            <h3><a href="{href}">{escape(title)}</a></h3>
            <p>{description}</p>
            <div class="essay-links"><a href="{href}">{read} <span aria-hidden="true">↗</span></a><a href="Articles/{safe}.docx" download>{download} <span aria-hidden="true">↓</span></a><span class="essay-filetype">HTML / DOCX</span></div>
          </div>
        </article>''')
        used.add(file.stem)
    if new_entries:
        marker = re.search(r'(?m)^      </div>\n      <p class="essay-note">', section)
        if not marker:
            raise ValueError("Cannot locate list ending for " + language)
        section = section[:marker.start()] + "\n".join(new_entries) + "\n" + section[marker.start():]

    # Ensure all entries have two functional links even if added manually.
    def normalize_entry(match: re.Match) -> str:
        entry = match.group(0)
        found = re.search(r'href="Articles/([^"#]+)\.html(?:#[a-z]+)?"', entry)
        if not found:
            raise ValueError("Essay entry has no HTML reading link: " + entry[:140])
        stem = found.group(1)
        if not (ARTICLES / (stem + ".html")).is_file() or not (ARTICLES / (stem + ".docx")).is_file():
            raise ValueError("Essay missing HTML/DOCX pair: " + stem)
        if not re.search(r'href="Articles/' + re.escape(stem) + r'\.docx"[^>]*download', entry):
            download = "Download" if language == "en" else "Скачать"
            link = f'<a href="Articles/{escape(stem, quote=True)}.docx" download>{download} <span aria-hidden="true">↓</span></a>'
            link_area = entry.find('<div class="essay-links">')
            if link_area == -1:
                raise ValueError("Missing essay-links for " + stem)
            entry = entry[:link_area+len('<div class="essay-links">')] + link + entry[link_area+len('<div class="essay-links">'):]
        entry = entry.replace('<span class="essay-filetype">HTML</span>', '<span class="essay-filetype">HTML / DOCX</span>')
        return entry
    section = re.sub(r'<article class="essay-entry">.*?</article>', normalize_entry, section, flags=re.S)

    num = 0
    def renumber(m: re.Match) -> str:
        nonlocal num
        num += 1
        return m.group(1) + str(num).zfill(2) + m.group(3)
    section = re.sub(r'(<div class="essay-number" aria-hidden="true">)(\d+)(</div>)',
                     renumber, section)
    section = re.sub(r'(<div class="essay-list-heading"><h3>[^<]+</h3><span>)[^<]*(</span>)',
                     lambda m: m.group(1) + (f"{num:02d} articles" if language=="en" else f"{num:02d} статей") + m.group(2), section)
    return before + section + after


def main() -> None:
    if not ARTICLES.is_dir():
        raise FileNotFoundError(ARTICLES)

    # Generate a missing counterpart only. Never replace an existing uploaded file.
    for source in sorted(ARTICLES.glob("*.docx")):
        output = source.with_suffix(".html")
        if not output.exists():
            docx_to_html(source, output)
    for source in sorted(ARTICLES.glob("*.html")):
        output = source.with_suffix(".docx")
        if not output.exists():
            html_to_docx(source, output)

    original = INDEX.read_text(encoding="utf-8")
    files = sorted(ARTICLES.glob("*.html"))
    changed = sync_listing(sync_listing(original, "en", files), "ru", files)
    if changed != original:
        INDEX.write_text(changed, encoding="utf-8")
        print("Updated bilingual essay catalog:", INDEX.relative_to(ROOT))
    else:
        print("Essay catalog already synchronized")


if __name__ == "__main__":
    main()
